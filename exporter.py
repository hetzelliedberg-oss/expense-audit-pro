import csv
import io
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXPORTS_FOLDER
from database import (
    get_transactions_filtered,
    get_daily_summary,
    get_category_breakdown,
    get_statistics
)

def export_transactions_to_excel(filters=None) -> str:
    """Generates an executive-grade styled multi-sheet Excel file (.xlsx)."""
    filters = filters or {}
    transactions = get_transactions_filtered(filters)
    daily_rows = get_daily_summary()
    cat_rows = get_category_breakdown()
    stats = get_statistics()

    wb = openpyxl.Workbook()

    # Style definitions
    font_family = "Cordia New"
    title_font = Font(name=font_family, size=18, bold=True, color="1F2937")
    sub_font = Font(name=font_family, size=13, color="4B5563")
    header_font = Font(name=font_family, size=14, bold=True, color="FFFFFF")
    cell_font = Font(name=font_family, size=13, color="111827")
    bold_cell_font = Font(name=font_family, size=13, bold=True, color="111827")
    warning_font = Font(name=font_family, size=13, bold=True, color="DC2626")

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy
    accent_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid") # Emerald
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    warn_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )

    # ----------------------------------------------------
    # SHEET 1: สรุปภาพรวม (Executive Summary)
    # ----------------------------------------------------
    ws_sum = wb.active
    ws_sum.title = "สรุปภาพรวม (Summary)"
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["B2"] = "รายงานตรวจสอบและวิเคราะห์ค่าใช้จ่าย (Expense Audit Report)"
    ws_sum["B2"].font = title_font
    ws_sum["B3"] = f"ออกรายงาน ณ วันที่: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_sum["B3"].font = sub_font

    # KPI Cards Block
    kpis = [
        ("จำนวนรายการทั้งหมด", f"{stats['total_count']:,} รายการ"),
        ("ยอดรวมรายจ่าย (Total Expense)", f"{stats['total_expense']:,.2f} ฿"),
        ("ยอดรวมรายรับ (Total Income)", f"{stats['total_income']:,.2f} ฿"),
        ("รายการที่พบข้อสังเกต/ซ้ำ (Flagged)", f"{stats['total_flagged']:,} รายการ")
    ]

    row_kpi = 5
    for idx, (label, val) in enumerate(kpis):
        r = row_kpi + idx
        ws_sum[f"B{r}"] = label
        ws_sum[f"B{r}"].font = bold_cell_font
        ws_sum[f"C{r}"] = val
        ws_sum[f"C{r}"].font = bold_cell_font
        ws_sum[f"C{r}"].alignment = Alignment(horizontal="right")
        ws_sum[f"B{r}"].fill = zebra_fill
        ws_sum[f"C{r}"].fill = zebra_fill

    # Category Breakdown Table
    cat_start_row = 11
    ws_sum[f"B{cat_start_row}"] = "สรุปค่าใช้จ่ายแยกตามหมวดหมู่"
    ws_sum[f"B{cat_start_row}"].font = Font(name=font_family, size=15, bold=True, color="1E3A8A")

    cat_headers = ["หมวดหมู่ค่าใช้จ่าย", "จำนวนรายการ", "ยอดเงินรวม (บาท)", "สัดส่วน (%)"]
    ws_sum.row_dimensions[cat_start_row + 1].height = 25
    for c_idx, h in enumerate(cat_headers, start=2):
        cell = ws_sum.cell(row=cat_start_row + 1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = accent_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cur_row = cat_start_row + 2
    total_exp = stats["total_expense"] if stats["total_expense"] > 0 else 1.0
    for c in cat_rows:
        pct = (c["total"] / total_exp) * 100.0
        ws_sum.cell(row=cur_row, column=2, value=c["category"]).font = cell_font
        ws_sum.cell(row=cur_row, column=3, value=c["count"]).font = cell_font
        cell_amt = ws_sum.cell(row=cur_row, column=4, value=c["total"])
        cell_amt.font = bold_cell_font
        cell_amt.number_format = "#,##0.00"
        cell_pct = ws_sum.cell(row=cur_row, column=5, value=f"{pct:.1f}%")
        cell_pct.font = cell_font
        cell_pct.alignment = Alignment(horizontal="right")
        cur_row += 1

    # ----------------------------------------------------
    # SHEET 2: สรุปรายวัน (Daily Breakdown)
    # ----------------------------------------------------
    ws_daily = wb.create_sheet(title="สรุปรายวัน (Daily)")
    ws_daily.views.sheetView[0].showGridLines = True
    ws_daily["A1"] = "สรุปค่าใช้จ่ายแยกแต่ละวัน (Daily Spending Breakdown)"
    ws_daily["A1"].font = title_font

    daily_headers = ["วันที่ (Date)", "จำนวนรายการ", "ยอดรายจ่ายรวม (บาท)", "ยอดรายรับรวม (บาท)", "รายการที่ต้องตรวจสอบ"]
    ws_daily.row_dimensions[3].height = 25
    for c_idx, h in enumerate(daily_headers, start=1):
        cell = ws_daily.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    d_row = 4
    for d in daily_rows:
        ws_daily.cell(row=d_row, column=1, value=d["transaction_date"]).font = cell_font
        ws_daily.cell(row=d_row, column=2, value=d["count"]).font = cell_font
        c_exp = ws_daily.cell(row=d_row, column=3, value=d["total_expense"])
        c_exp.font = bold_cell_font
        c_exp.number_format = "#,##0.00"
        c_inc = ws_daily.cell(row=d_row, column=4, value=d["total_income"])
        c_inc.font = cell_font
        c_inc.number_format = "#,##0.00"
        c_warn = ws_daily.cell(row=d_row, column=5, value=d["warning_count"])
        c_warn.font = warning_font if d["warning_count"] > 0 else cell_font
        c_warn.alignment = Alignment(horizontal="center")
        d_row += 1

    # ----------------------------------------------------
    # SHEET 3: รายการทั้งหมด (All Transactions)
    # ----------------------------------------------------
    ws_tx = wb.create_sheet(title="รายการทั้งหมด (Transactions)")
    ws_tx.views.sheetView[0].showGridLines = True
    ws_tx["A1"] = "รายละเอียดธุรกรรมทั้งหมด (Transaction Audit Log)"
    ws_tx["A1"].font = title_font

    tx_headers = [
        "รหัส (ID)", "วันที่ (Date)", "เวลา (Time)", "ประเภท", "ยอดเงิน (Amount)",
        "ค่าธรรมเนียม", "ภาษี VAT", "ยอดรวมสุทธิ", "หมวดหมู่", "หมวดหมู่ย่อย",
        "ช่องทาง/ธนาคาร", "ผู้รับเงิน / ร้านค้า", "ผู้โอนเงิน", "รหัสอ้างอิงสลิป",
        "สถานะ Audit", "ข้อสังเกตจากการตรวจสอบ", "รายการสินค้า/บันทึก"
    ]

    ws_tx.row_dimensions[3].height = 25
    for c_idx, h in enumerate(tx_headers, start=1):
        cell = ws_tx.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    t_row = 4
    for tx in transactions:
        is_warn = tx["audit_status"] != "clean"
        ws_tx.cell(row=t_row, column=1, value=tx["id"]).font = cell_font
        ws_tx.cell(row=t_row, column=2, value=tx["transaction_date"]).font = cell_font
        ws_tx.cell(row=t_row, column=3, value=tx["transaction_time"]).font = cell_font
        ws_tx.cell(row=t_row, column=4, value="รายจ่าย" if tx["type"] == "expense" else ("รายรับ" if tx["type"] == "income" else "โอน")).font = cell_font
        
        c_amt = ws_tx.cell(row=t_row, column=5, value=tx["amount"])
        c_amt.font = cell_font
        c_amt.number_format = "#,##0.00"

        c_fee = ws_tx.cell(row=t_row, column=6, value=tx["fee"])
        c_fee.font = cell_font
        c_fee.number_format = "#,##0.00"

        c_vat = ws_tx.cell(row=t_row, column=7, value=tx["vat"])
        c_vat.font = cell_font
        c_vat.number_format = "#,##0.00"

        c_tot = ws_tx.cell(row=t_row, column=8, value=tx["total_amount"])
        c_tot.font = bold_cell_font
        c_tot.number_format = "#,##0.00"

        ws_tx.cell(row=t_row, column=9, value=tx["category"]).font = cell_font
        ws_tx.cell(row=t_row, column=10, value=tx["subcategory"]).font = cell_font
        ws_tx.cell(row=t_row, column=11, value=tx["payment_source"]).font = cell_font
        ws_tx.cell(row=t_row, column=12, value=tx["payee_name"]).font = cell_font
        ws_tx.cell(row=t_row, column=13, value=tx["sender_name"]).font = cell_font
        ws_tx.cell(row=t_row, column=14, value=tx["ref_number"]).font = cell_font
        
        c_stat = ws_tx.cell(row=t_row, column=15, value=tx["audit_status"])
        c_stat.font = warning_font if is_warn else cell_font
        if is_warn:
            c_stat.fill = warn_fill

        c_note = ws_tx.cell(row=t_row, column=16, value=tx["audit_note"])
        c_note.font = warning_font if is_warn else cell_font

        ws_tx.cell(row=t_row, column=17, value=tx["notes"]).font = cell_font
        t_row += 1

    # Auto-adjust column widths on all worksheets
    for ws in [ws_sum, ws_daily, ws_tx]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"expense_audit_report_{timestamp}.xlsx"
    file_path = EXPORTS_FOLDER / filename
    wb.save(str(file_path))

    return str(file_path)

def export_transactions_to_csv(filters=None) -> str:
    """Exports transactions to a CSV file encoded in UTF-8 with BOM for Thai Excel."""
    filters = filters or {}
    transactions = get_transactions_filtered(filters)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"expense_audit_report_{timestamp}.csv"
    file_path = EXPORTS_FOLDER / filename

    headers = [
        "ID", "วันที่", "เวลา", "ประเภท", "ยอดเงิน", "ค่าธรรมเนียม", "ภาษีVAT",
        "ยอดรวมสุทธิ", "หมวดหมู่", "หมวดหมู่ย่อย", "ช่องทางธนาคาร", "ผู้รับเงิน/ร้านค้า",
        "ผู้โอนเงิน", "รหัสอ้างอิงสลิป", "สถานะAudit", "ข้อสังเกตการตรวจสอบ", "บันทึกช่วยจำ"
    ]

    with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for tx in transactions:
            writer.writerow([
                tx["id"],
                tx["transaction_date"],
                tx["transaction_time"],
                "รายจ่าย" if tx["type"] == "expense" else ("รายรับ" if tx["type"] == "income" else "โอน"),
                f"{tx['amount']:.2f}",
                f"{tx['fee']:.2f}",
                f"{tx['vat']:.2f}",
                f"{tx['total_amount']:.2f}",
                tx["category"],
                tx["subcategory"],
                tx["payment_source"],
                tx["payee_name"],
                tx["sender_name"],
                tx["ref_number"],
                tx["audit_status"],
                tx["audit_note"],
                tx["notes"]
            ])

    return str(file_path)
